"""
文档解析引擎
支持 DOCX, XLSX, PDF 格式的混合格式提取
输出统一的 _parsed.json 结构化数据
"""

import os
import json
import logging
from datetime import datetime
from typing import List, Dict, Any, Optional

from app.utils.logger import get_logger


class DocumentParser:
    """文档解析器基类

    支持三种文档格式的解析:
    - DOCX: Word 文档，提取段落、标题和表格
    - XLSX: Excel 工作簿，提取工作表和表格数据
    - PDF: PDF 文档，提取文本和表格

    所有格式统一输出为结构化 JSON 数据，包含 sections 和 plain_text_preview
    """

    def __init__(self):
        self.supported_formats = ['.docx', '.xlsx', '.pdf']
        self.logger = get_logger("document_parser")

    def parse_file(self, file_path: str) -> Dict[str, Any]:
        """解析文档文件并返回结构化数据

        Args:
            file_path: 原始文件路径

        Returns:
            解析后的结构化字典，包含:
            - doc_id: 文档ID (空字符串，由调用方填充)
            - filename: 文件名
            - file_type: 文件类型 (docx/xlsx/pdf)
            - parsed_at: 解析时间戳
            - total_length: 文本总长度
            - sections: 结构化内容列表
            - plain_text_preview: 纯文本预览

        Raises:
            FileNotFoundError: 文件不存在
            ValueError: 不支持的文件格式或解析失败
        """
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"文件不存在: {file_path}")

        ext = os.path.splitext(file_path)[1].lower()
        self.logger.info(f"开始解析文件: {file_path} (类型: {ext})")

        try:
            if ext == '.docx':
                result = self._parse_docx(file_path)
            elif ext == '.xlsx':
                result = self._parse_xlsx(file_path)
            elif ext == '.pdf':
                result = self._parse_pdf(file_path)
            else:
                raise ValueError(f"不支持的文件格式: {ext}，支持格式: {self.supported_formats}")

            self.logger.info(f"文件解析成功: {file_path}, 总长度: {result.get('total_length', 0)}")
            return result

        except ImportError as e:
            error_msg = f"缺少必要的依赖库: {e}"
            self.logger.error(error_msg)
            return {
                'error': error_msg,
                'filename': os.path.basename(file_path),
                'file_type': ext.lstrip('.'),
                'parsed_at': datetime.now().isoformat(),
                'sections': [],
                'plain_text_preview': ''
            }
        except Exception as e:
            error_msg = f"文件解析失败: {str(e)}"
            self.logger.error(f"{error_msg} (文件: {file_path})")
            raise ValueError(error_msg)

    def _parse_docx(self, file_path: str) -> Dict[str, Any]:
        """解析 DOCX 文件

        提取内容:
        - 标题段落 (Heading 1-3)
        - 普通段落
        - 表格 (转为 Markdown 格式)

        Args:
            file_path: DOCX 文件路径

        Returns:
            dict: 解析结果
        """
        try:
            from docx import Document
        except ImportError:
            raise ImportError("python-docx 库未安装，请执行: pip install python-docx")

        try:
            doc = Document(file_path)
        except Exception as e:
            raise ValueError(f"无法读取 DOCX 文件: {e}")

        sections = []
        plain_text_parts = []

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            style_name = para.style.name if para.style else ''

            if style_name.startswith('Heading'):
                level = int(''.join(filter(str.isdigit, style_name)) or '1')
                level = min(max(level, 1), 3)

                sections.append({
                    'type': 'heading',
                    'level': level,
                    'content': text
                })
                plain_text_parts.append(f"\n{'#' * level} {text}\n")
            else:
                sections.append({
                    'type': 'paragraph',
                    'content': text
                })
                plain_text_parts.append(text + '\n')

        for table in doc.tables:
            headers = [cell.text.strip() for cell in table.rows[0].cells]
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows[1:]]

            markdown_table = self._build_markdown_table(headers, rows)

            sections.append({
                'type': 'table',
                'headers': headers,
                'rows': rows,
                'markdown': markdown_table
            })
            plain_text_parts.append('\n' + markdown_table + '\n')

        return {
            'doc_id': '',
            'filename': os.path.basename(file_path),
            'file_type': 'docx',
            'parsed_at': datetime.now().isoformat(),
            'total_length': sum(len(p) for p in plain_text_parts),
            'sections': sections,
            'plain_text_preview': ''.join(plain_text_parts).strip()
        }

    def _parse_xlsx(self, file_path: str) -> Dict[str, Any]:
        """解析 XLSX 文件

        提取内容:
        - 每个工作表作为 H2 标题
        - 表格数据 (第一行作为表头)

        Args:
            file_path: XLSX 文件路径

        Returns:
            dict: 解析结果
        """
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise ImportError("openpyxl 库未安装，请执行: pip install openpyxl")

        try:
            wb = load_workbook(file_path, read_only=True, data_only=True)
        except Exception as e:
            raise ValueError(f"无法读取 XLSX 文件: {e}")

        sections = []
        plain_text_parts = []

        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]

            sections.append({
                'type': 'heading',
                'level': 2,
                'content': f"工作表: {sheet_name}"
            })
            plain_text_parts.append(f"\n## 工作表: {sheet_name}\n\n")

            data_rows = []
            for i, row in enumerate(sheet.iter_rows(values_only=True)):
                row_data = [str(cell) if cell is not None else '' for cell in row]

                if all(not cell for cell in row_data):
                    continue

                data_rows.append(row_data)

            if data_rows:
                headers = data_rows[0]
                rows = data_rows[1:] if len(data_rows) > 1 else []

                if headers or rows:
                    markdown_table = self._build_markdown_table(headers, rows)
                    sections.append({
                        'type': 'table',
                        'headers': headers,
                        'rows': rows,
                        'markdown': markdown_table
                    })
                    plain_text_parts.append('\n' + markdown_table + '\n')

        wb.close()

        return {
            'doc_id': '',
            'filename': os.path.basename(file_path),
            'file_type': 'xlsx',
            'parsed_at': datetime.now().isoformat(),
            'total_length': sum(len(p) for p in plain_text_parts),
            'sections': sections,
            'plain_text_preview': ''.join(plain_text_parts).strip()
        }

    def _parse_pdf(self, file_path: str) -> Dict[str, Any]:
        """解析 PDF 文件

        提取内容:
        - 每页文本内容 (带页码标记)
        - 页面中的表格 (转为 Markdown 格式)

        Args:
            file_path: PDF 文件路径

        Returns:
            dict: 解析结果
        """
        try:
            import pdfplumber
        except ImportError:
            raise ImportError("pdfplumber 库未安装，请执行: pip install pdfplumber")

        sections = []
        plain_text_parts = []

        try:
            with pdfplumber.open(file_path) as pdf:
                total_pages = len(pdf.pages)

                for page_num, page in enumerate(pdf.pages, 1):
                    text = page.extract_text() or ''
                    if text.strip():
                        sections.append({
                            'type': 'paragraph',
                            'content': text.strip(),
                            'metadata': {'page': page_num}
                        })
                        plain_text_parts.append(f'\n--- 第 {page_num}/{total_pages} 页 ---\n\n{text}\n')

                    tables = page.extract_tables()
                    for table_idx, table in enumerate(tables):
                        if table and len(table) > 0:
                            headers = [str(cell).strip() if cell else '' for cell in table[0]]
                            rows = [[str(cell).strip() if cell else '' for cell in row] for row in table[1:]]

                            if headers or rows:
                                markdown_table = self._build_markdown_table(headers, rows)
                                sections.append({
                                    'type': 'table',
                                    'headers': headers,
                                    'rows': rows,
                                    'markdown': markdown_table,
                                    'metadata': {'page': page_num, 'table_index': table_idx}
                                })
                                plain_text_parts.append('\n' + markdown_table + '\n')
        except Exception as e:
            raise ValueError(f"无法读取 PDF 文件: {e}")

        return {
            'doc_id': '',
            'filename': os.path.basename(file_path),
            'file_type': 'pdf',
            'parsed_at': datetime.now().isoformat(),
            'total_length': sum(len(p) for p in plain_text_parts),
            'sections': sections,
            'plain_text_preview': ''.join(plain_text_parts).strip()
        }

    def _build_markdown_table(self, headers: List[str], rows: List[List[str]]) -> str:
        """构建 Markdown 表格字符串

        Args:
            headers: 表头列表
            rows: 数据行列表

        Returns:
            str: Markdown 格式的表格字符串
        """
        if not headers and not rows:
            return ''

        lines = []
        if headers:
            lines.append('| ' + ' | '.join(headers) + ' |')
            lines.append('|' + '|'.join(['---'] * len(headers)) + '|')

        for row in rows:
            max_cols = len(headers) if headers else len(row)
            lines.append('| ' + ' | '.join(row[:max_cols]) + ' |')

        return '\n'.join(lines)

    def save_parsed_data(self, parsed_data: Dict, output_path: str) -> bool:
        """保存解析结果到 JSON 文件

        Args:
            parsed_data: 解析后的数据字典
            output_path: 输出文件路径

        Returns:
            bool: 保存是否成功
        """
        try:
            output_dir = os.path.dirname(output_path)
            if output_dir and not os.path.exists(output_dir):
                os.makedirs(output_dir, exist_ok=True)

            with open(output_path, 'w', encoding='utf-8') as f:
                json.dump(parsed_data, f, ensure_ascii=False, indent=2)

            self.logger.debug(f"解析结果已保存: {output_path}")
            return True

        except IOError as e:
            self.logger.error(f"保存解析结果失败: {output_path}, 错误: {e}")
            return False


document_parser = DocumentParser()
